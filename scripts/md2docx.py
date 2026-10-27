#!/usr/bin/env python
"""简易 markdown -> docx(简历版式,已固定页边距与字号)。

用法:
  python md2docx.py <输入.md> [输出.docx]

支持的标记(只这些,别塞复杂语法):
  # 标题        -> 姓名,20pt 深蓝居中
  ## 小节       -> 12pt 深蓝加粗 + 下边框
  - 条目        -> 「• 」项目符号
    - 子条目    -> 缩进「– 」
  **粗体**      -> 行内粗体
  第一个 ## 之前的非空行 -> 抬头联系行(9.5pt 灰居中)
"""
import sys
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = '微软雅黑'
BODY = Pt(10)
ACCENT = RGBColor(0x1F, 0x3B, 0x63)
GRAY = RGBColor(0x59, 0x59, 0x59)


def set_font(run, size=BODY, bold=False, color=None):
    run.font.name = FONT
    run.font.size = size
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia'):
        rfonts.set(qn(attr), FONT)


def add_runs(p, text, size=BODY, color=None, base_bold=False):
    for i, seg in enumerate(re.split(r'\*\*(.+?)\*\*', text)):
        if not seg:
            continue
        bold = base_bold or (i % 2 == 1)
        set_font(p.add_run(seg), size=size, bold=bold, color=color)


def bottom_border(p):
    ppr = p._p.get_or_add_pPr()
    borders = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '8')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '1F3B63')
    borders.append(bottom)
    ppr.append(borders)


def main():
    if len(sys.argv) < 2:
        print('用法: python md2docx.py <输入.md> [输出.docx]')
        return 2
    src = sys.argv[1]
    dst = sys.argv[2] if len(sys.argv) > 2 else os.path.splitext(src)[0] + '.docx'

    with open(src, encoding='utf-8') as f:
        lines = f.read().splitlines()

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.6)
    sec.bottom_margin = Cm(1.6)
    sec.left_margin = Cm(1.9)
    sec.right_margin = Cm(1.9)

    normal = doc.styles['Normal']
    normal.font.name = FONT
    normal.font.size = BODY
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

    leading = True  # 第一个 ## 之前当抬头区

    for raw in lines:
        line = raw.rstrip()
        if not line.strip():
            continue
        stripped = line.strip()

        if stripped.startswith('# '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, stripped[2:], size=Pt(20), color=ACCENT, base_bold=True)
            continue

        if stripped.startswith('## '):
            leading = False
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            add_runs(p, stripped[3:], size=Pt(12), color=ACCENT, base_bold=True)
            bottom_border(p)
            continue

        if leading:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            add_runs(p, stripped, size=Pt(9.5), color=GRAY)
            continue

        m = re.match(r'^(\s*)-\s+(.*)$', line)
        if m:
            depth = len(m.group(1)) // 2
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Cm(0.55 + 0.55 * depth)
            pf.first_line_indent = Cm(-0.45)
            pf.space_after = Pt(1.5)
            pf.line_spacing = 1.2
            set_font(p.add_run('• ' if depth == 0 else '– '), size=BODY, color=ACCENT)
            add_runs(p, m.group(2))
            continue

        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(2)
        pf.line_spacing = 1.2
        add_runs(p, stripped)

    doc.save(dst)
    print('OK', dst)
    return 0


if __name__ == '__main__':
    sys.exit(main())
